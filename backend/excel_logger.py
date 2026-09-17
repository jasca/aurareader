"""
Módulo para registro automático de atenciones/consultantes de AuraReader en un archivo Excel.
Ubicación: backend/consultantes_aurareader.xlsx
"""

import os
import logging
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

EXCEL_FILE_PATH = os.path.abspath(os.path.join(os.path.dirname(__file__), "consultantes_aurareader.xlsx"))

def log_scan_to_excel(servicio: str, nombre: str, telefono: str, email: str, resultado: str, pdf_filename: str):
    try:
        file_exists = os.path.exists(EXCEL_FILE_PATH)
        
        if file_exists:
            wb = openpyxl.load_workbook(EXCEL_FILE_PATH)
            ws = wb.active
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Consultantes AuraReader"
            
            # Encabezados con estilo
            headers = [
                "Fecha y Hora", 
                "Servicio", 
                "Nombre y Apellido", 
                "Teléfono", 
                "Email", 
                "Resultado / Frecuencia", 
                "Archivo PDF Generado"
            ]
            ws.append(headers)
            
            header_fill = PatternFill(start_color="7C3AED", end_color="7C3AED", fill_type="solid") # Violeta holístico
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            thin_border = Border(
                left=Side(style='thin', color='DDDDDD'),
                right=Side(style='thin', color='DDDDDD'),
                top=Side(style='thin', color='DDDDDD'),
                bottom=Side(style='thin', color='DDDDDD')
            )
            
            for col_num, header_text in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border
                
            ws.row_dimensions[1].height = 25

        # Agregar nueva fila con los datos de la atención
        fecha_hora = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
        row_data = [
            fecha_hora,
            servicio,
            nombre.strip() if nombre else "Sin Nombre",
            telefono.strip() if telefono else "-",
            email.strip() if email else "-",
            resultado,
            pdf_filename
        ]
        ws.append(row_data)
        
        current_row = ws.max_row
        row_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )
        
        for col_num in range(1, len(row_data) + 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.border = row_border
            cell.font = Font(name="Calibri", size=11)
            if col_num in [1, 2, 4]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

        # Ajustar ancho de columnas automáticamente
        column_widths = [20, 15, 25, 18, 28, 45, 30]
        for idx, width in enumerate(column_widths, 1):
            col_letter = openpyxl.utils.get_column_letter(idx)
            ws.column_dimensions[col_letter].width = width

        wb.save(EXCEL_FILE_PATH)
        logging.info(f"Fila registrada exitosamente en Excel: {EXCEL_FILE_PATH}")
    except Exception as e:
        logging.error(f"Error al guardar datos en Excel: {e}")
